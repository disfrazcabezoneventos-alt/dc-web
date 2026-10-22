"""Build the reviewed, public-only Cotillón import from the local workbook.

Run locally with openpyxl; the source workbook and hidden products stay private.
"""
import json
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\Users\Usuario\OneDrive\Escritorio\Disfraz_Cabezon_Catalogo.xlsx")
OUTPUT = Path("backend/import-cotillon-excel-2026.json")
HEADERS = ["nombre", "descripcion", "categoria_principal", "subcategoria", "subcategoria_interna", "precio", "stock", "foto", "visible"]

book = openpyxl.load_workbook(SOURCE, data_only=True)
web = book["Web dc"]
general = book["General"]
assert [cell.value for cell in next(web.iter_rows(min_row=1, max_row=1))] == HEADERS
assert web.max_row == 1207 and general.max_row == 1212

products = []
hidden = 0
for row_number in range(2, web.max_row + 1):
    data = dict(zip(HEADERS, (web.cell(row_number, column).value for column in range(1, 10))))
    general_row = row_number + 5
    assert data["foto"] == general.cell(general_row, 18).value
    assert data["visible"] in ("Sí", "No")
    if data["visible"] == "No":
        hidden += 1
        continue
    group, section, inner = data["categoria_principal"], data["subcategoria"] or "", data["subcategoria_interna"] or ""
    if row_number == 398:
        assert data["nombre"] == "Cortante Momia"
        group, section, inner = "REPOSTERÍA", "Accesorios", ""
    if row_number == 723:
        assert data["nombre"] == "Mantel Friselina Colores Varios Rectangular"
        group, section, inner = "DESCARTABLES", "", ""
    source_id = str(general.cell(general_row, 6).value).strip()
    barcode = str(general.cell(general_row, 5).value).strip()
    source_image = str(general.cell(general_row, 19).value).strip()
    assert source_id and barcode, (row_number, source_id, barcode)
    assert source_image.startswith("https://images.rappi.com.ar/products/")
    assert isinstance(data["precio"], (int, float)) and data["precio"] > 0
    assert isinstance(data["stock"], int) and data["stock"] > 0
    products.append({
        "sourceId": source_id,
        "sku": barcode,
        "name": data["nombre"].strip(),
        "description": data["descripcion"].strip(),
        "group": group,
        "section": section,
        "inner": inner,
        "price": data["precio"],
        "stock": data["stock"],
        "sourceImage": source_image,
        "photoReview": row_number == 398,
        "excelRow": row_number,
    })

assert len(products) == 790 and hidden == 416
assert len({item["sourceId"] for item in products}) == len(products)
assert len({item["sku"] for item in products}) == len(products)
assert len({item["sourceImage"] for item in products}) == len(products)
assert len([item for item in products if not item["section"]]) == 1
OUTPUT.write_text(json.dumps({"version": 1, "source": "Disfraz_Cabezon_Catalogo.xlsx / Web dc", "products": products}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"Exportados {len(products)} visibles. {hidden} ocultos permanecen sólo en el Excel. Archivo: {OUTPUT}")
