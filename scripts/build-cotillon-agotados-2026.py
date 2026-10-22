"""Export the 416 inactive/out-of-stock rows without changing the first import."""
import json
from pathlib import Path

import openpyxl

source = Path(r"C:\Users\Usuario\OneDrive\Escritorio\Disfraz_Cabezon_Catalogo.xlsx")
output = Path("backend/import-cotillon-agotados-2026.json")
book = openpyxl.load_workbook(source, data_only=True)
web, general = book["Web dc"], book["General"]
headers = ["nombre", "descripcion", "categoria_principal", "subcategoria", "subcategoria_interna", "precio", "stock", "foto", "visible"]
assert [cell.value for cell in web[1]] == headers
assert web.max_row == 1207 and general.max_row == 1212
products = []
for number in range(2, web.max_row + 1):
    row = dict(zip(headers, (web.cell(number, column).value for column in range(1, 10))))
    assert row["visible"] in ("Sí", "No")
    if row["visible"] != "No":
        continue
    original = number + 5
    assert row["foto"] == general.cell(original, 18).value
    assert row["stock"] is None or row["stock"] == 0
    url = str(general.cell(original, 19).value or "").strip()
    assert url.startswith("https://images.rappi.com.ar/products/")
    products.append({
        "sourceId": str(general.cell(original, 6).value or "").strip(),
        "sku": str(general.cell(original, 5).value or "").strip(),
        "name": row["nombre"].strip(),
        "description": row["descripcion"].strip(),
        "group": row["categoria_principal"],
        "section": row["subcategoria"] or "",
        "inner": row["subcategoria_interna"] or "",
        "price": row["precio"],
        "stock": 0,
        "sourceImage": url,
        "excelRow": number,
    })
assert len(products) == 416
assert len({item["sourceId"] for item in products}) == 416
assert len({item["sourceImage"] for item in products}) == 416
assert [item["name"] for item in products if not item["section"]] == ["Bowl Blanco"]
output.write_text(json.dumps({"version": 1, "source": "Disfraz_Cabezon_Catalogo.xlsx / Web dc / visible=No", "products": products}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"Exportados {len(products)} productos agotados a {output}")
